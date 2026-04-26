"""
Generate SRS Document in DOCX — IEEE 830 Template Format (Karl Wiegers)
Traffic Prediction and Analysis — Group E-16
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
import os

doc = Document()

# ============ STYLES ============
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Times New Roman'
    hs.font.color.rgb = RGBColor(0, 0, 0)

def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def add_tbl(doc, headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = h
        for p in c.paragraphs:
            for r in p.runs: r.bold = True; r.font.size = Pt(10)
        shade(c, "D9E2F3")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri+1].cells[ci]; c.text = str(val)
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(10)
    return t

def bp(doc, items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

def add_feature(doc, num, name, priority, description, stimuli, reqs):
    """Add a System Feature section per IEEE template format"""
    doc.add_heading(f"4.{num} {name}", level=2)
    doc.add_heading(f"4.{num}.1 Description and Priority", level=3)
    p = doc.add_paragraph()
    r = p.add_run(f"Priority: {priority}"); r.bold = True
    doc.add_paragraph(description)
    doc.add_heading(f"4.{num}.2 Stimulus/Response Sequences", level=3)
    for s in stimuli:
        doc.add_paragraph(s, style='List Bullet')
    doc.add_heading(f"4.{num}.3 Functional Requirements", level=3)
    for rid, req in reqs:
        p = doc.add_paragraph()
        r = p.add_run(f"{rid}: "); r.bold = True
        p.add_run(req)

# ============ TITLE PAGE (matching IEEE template) ============
for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Software Requirements\nSpecification"); r.bold = True; r.font.size = Pt(28)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("for"); r.font.size = Pt(16)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Traffic Prediction and Analysis"); r.bold = True; r.font.size = Pt(22)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Version 1.0 approved"); r.font.size = Pt(14)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prepared by\nHarshal Jaiswal, Pranav Jagtap, Suyash Jobanputra, Shreyash Jobanputra\n(Group E-16)"); r.font.size = Pt(12)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Vishwakarma Institute of Technology, Pune\nDepartment of Computer Engineering"); r.font.size = Pt(12)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("March 2, 2026"); r.font.size = Pt(12)
doc.add_page_break()

# ============ TABLE OF CONTENTS (IEEE format) ============
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Software Requirements Specification for Traffic Prediction and Analysis"); r.italic = True; r.font.size = Pt(10)
doc.add_paragraph()
doc.add_heading("Table of Contents", level=1)
toc = [
    "Revision History",
    "1.  Introduction",
    "    1.1  Purpose",
    "    1.2  Document Conventions",
    "    1.3  Intended Audience and Reading Suggestions",
    "    1.4  Product Scope",
    "    1.5  References",
    "2.  Overall Description",
    "    2.1  Product Perspective",
    "    2.2  Product Functions",
    "    2.3  User Classes and Characteristics",
    "    2.4  Operating Environment",
    "    2.5  Design and Implementation Constraints",
    "    2.6  User Documentation",
    "    2.7  Assumptions and Dependencies",
    "3.  External Interface Requirements",
    "    3.1  User Interfaces",
    "    3.2  Hardware Interfaces",
    "    3.3  Software Interfaces",
    "    3.4  Communications Interfaces",
    "4.  System Features",
    "    4.1  User Authentication and Management",
    "    4.2  Traffic Prediction Engine",
    "    4.3  Pothole Detection (Machine Learning)",
    "    4.4  Citizen Complaint Management",
    "    4.5  Event and Infrastructure Management",
    "    4.6  Traffic Analytics and Reporting",
    "    4.7  Automated Data Collection",
    "    4.8  Map Visualization and Location Services",
    "    4.9  Admin Dashboard",
    "    4.10  Prediction Accuracy and Logging",
    "    4.11  Citizen Rewards",
    "5.  Other Nonfunctional Requirements",
    "    5.1  Performance Requirements",
    "    5.2  Safety Requirements",
    "    5.3  Security Requirements",
    "    5.4  Software Quality Attributes",
    "    5.5  Business Rules",
    "6.  Other Requirements",
    "Appendix A: Glossary",
    "Appendix B: Analysis Models",
    "Appendix C: Technology Stack",
    "Appendix D: API Specification",
    "Appendix E: Data Dictionary",
]
for item in toc:
    p = doc.add_paragraph(item); p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(0)
doc.add_paragraph()

# Revision History
doc.add_heading("Revision History", level=1)
add_tbl(doc, ["Name", "Date", "Reason For Changes", "Version"], [
    ["Group E-16", "Oct 2025", "Initial draft with core requirements", "0.1"],
    ["Group E-16", "Dec 2025", "Added ML, analytics, admin requirements", "0.5"],
    ["Group E-16", "Feb 2026", "Final version with all implemented features", "1.0"],
])
doc.add_page_break()

# ============================================================
# 1. INTRODUCTION
# ============================================================
doc.add_heading("1. Introduction", level=1)

doc.add_heading("1.1 Purpose", level=2)
doc.add_paragraph(
    "This document specifies the software requirements for the Traffic Prediction and Analysis system, "
    "Version 1.0. The system is a web-based platform developed for Pune Municipal Corporation (PMC) "
    "that enables real-time traffic prediction, citizen complaint management, AI-powered pothole "
    "detection, and comprehensive traffic analytics for the city of Pune."
)
doc.add_paragraph(
    "This SRS covers the complete system comprising a React frontend, Node.js Express backend, "
    "and Flask ML microservice."
)

doc.add_heading("1.2 Document Conventions", level=2)
doc.add_paragraph("This document follows these typographical conventions:")
bp(doc, [
    "Bold text indicates requirement IDs (e.g., REQ-1, REQ-2).",
    "SHALL indicates a mandatory requirement implemented in the current version.",
    "SHOULD indicates a recommended feature for future enhancement.",
    "Higher-level feature priorities are inherited by all child functional requirements unless explicitly overridden.",
    "All requirement IDs are unique across the system and follow the format REQ-<feature>-<number>.",
])

doc.add_heading("1.3 Intended Audience and Reading Suggestions", level=2)
doc.add_paragraph("This document is intended for the following audiences:")
add_tbl(doc, ["Audience", "Suggested Reading"], [
    ["PMC Officials / Stakeholders", "Sections 1, 2, 4 (System Features overview), and 5"],
    ["Development Team (Group E-16)", "All sections, especially 3, 4, and Appendices C-E"],
    ["Project Guide (Dr. Radhika Kulkarni)", "Complete document for academic evaluation"],
    ["Testers / QA", "Sections 4 (functional requirements), 5 (nonfunctional), and Appendix D (APIs)"],
    ["System Administrators", "Sections 2.4, 2.5, 3, 5.3, and Appendix C"],
])
doc.add_paragraph(
    "Readers should begin with Sections 1 and 2 for project overview, then proceed to Section 4 "
    "for detailed feature requirements. Section 3 details external interfaces, Section 5 covers "
    "nonfunctional requirements, and the Appendices provide technical reference material."
)

doc.add_heading("1.4 Product Scope", level=2)
doc.add_paragraph(
    "The Traffic Prediction and Analysis system is a Major Project (Group E-16, AY 2025-2026) "
    "developed at Vishwakarma Institute of Technology, Pune, in collaboration with Pune Municipal Corporation. "
    "The software provides intelligent traffic congestion prediction using a multi-factor weighted "
    "scoring algorithm, citizen engagement through complaint and pothole reporting, and comprehensive "
    "analytics dashboards for municipal traffic management."
)
doc.add_paragraph("Key benefits and objectives:")
bp(doc, [
    "Enable citizens to report road issues (potholes, complaints) with AI-verified image uploads",
    "Provide multi-factor traffic prediction using historical data, weather, events, Google Traffic, and more",
    "Offer route analysis incorporating nearby schools, hospitals, hotspots, and construction zones",
    "Calculate a traffic score (0-100) for route safety and efficiency assessment",
    "Supply analytics dashboards with trends, heatmaps, and root-cause analysis for PMC officials",
    "Automate scheduled traffic data collection for monitored Pune routes",
])

doc.add_heading("1.5 References", level=2)
refs = [
    "IEEE 830-1998 — IEEE Recommended Practice for Software Requirements Specifications",
    "B.Tech-14 Major Project Blackbook — Data Driven PMC Traffic Analysis System",
    "Ahmed, K.R. & Kharel, S. (2021). Potholes Detection Using Deep Learning. DOI: 10.1007/978-3-030-82199-9_24",
    "Grepon, B. et al. (2023). RUI: Web-based Road Updates Information System using Google Maps API. IJCSR 7, 2253-2271",
    "Lin, P. et al. (2022). Data-driven spatial-temporal analysis of highway traffic volume. TBS 29, 95-112",
    "Medina-Salgado, B. et al. (2022). Urban traffic flow prediction techniques: A review. Sustainable Computing 35",
    "MongoDB Documentation — https://www.mongodb.com/docs/",
    "Express.js Documentation — https://expressjs.com/",
    "React Documentation — https://react.dev/",
    "TensorFlow/Keras MobileNet — https://keras.io/api/applications/mobilenet/",
    "OpenWeatherMap API — https://openweathermap.org/api",
    "Google Maps Directions API — https://developers.google.com/maps/documentation/directions",
]
for i, ref in enumerate(refs, 1):
    doc.add_paragraph(f"[{i}] {ref}")
doc.add_page_break()

# ============================================================
# 2. OVERALL DESCRIPTION
# ============================================================
doc.add_heading("2. Overall Description", level=1)

doc.add_heading("2.1 Product Perspective", level=2)
doc.add_paragraph(
    "The Traffic Prediction and Analysis system is a new, self-contained web application designed "
    "for PMC's traffic management department. It is not a replacement or follow-on of any existing system. "
    "The system operates as a three-tier client-server architecture:"
)
bp(doc, [
    "React + Vite Frontend (Port 5173) — Single-page web application with dual-map interface",
    "Node.js Express Backend (Port 3001) — RESTful API server with 22 controllers, 19 models, 20 route modules",
    "Flask ML Microservice (Port 3003) — MobileNet-based pothole image classification service",
])
doc.add_paragraph(
    "The system interfaces with external services: OpenWeatherMap API (weather data), Google Maps "
    "Directions API (traffic durations), OpenStreetMap (map tiles), Cloudinary (image storage), "
    "and MongoDB Atlas (data persistence with geospatial indexing)."
)

doc.add_heading("2.2 Product Functions", level=2)
doc.add_paragraph("The major functions the product performs are:")
bp(doc, [
    "User Authentication — Registration, login/logout with JWT, role-based access (User/Admin)",
    "Traffic Prediction — Multi-factor weighted scoring (0-100) using 7+ data sources",
    "Smart Time Suggestions — Recommend optimal travel times by comparing adjacent time slots",
    "Pothole Detection — AI-powered MobileNet image classification (pothole/no-pothole)",
    "Citizen Complaints — Submit geotagged complaints with image evidence",
    "Event & Infrastructure Management — Track events, schools, hospitals, construction, diversions",
    "Traffic Analytics — Trends, heatmaps, root-cause breakdowns, route comparisons, AI recommendations",
    "Automated Data Collection — Scheduled 2-hourly collection for 5 monitored Pune routes",
    "Admin Dashboard — Manage users, complaints, potholes, analytics, prediction accuracy",
    "Map Visualization — Dual-map interface with Leaflet and Google Maps",
    "Prediction Accuracy Logging — Track predicted vs. actual scores over time",
    "Citizen Rewards — Gamification for issue reporting contributions",
])

doc.add_heading("2.3 User Classes and Characteristics", level=2)
add_tbl(doc, ["User Class", "Description", "Technical Expertise", "Primary Functions Used"], [
    ["Citizen (Registered)", "General public, daily commuters", "Low — uses web browser interface", "Traffic prediction, complaint/pothole reporting, route analysis, rewards"],
    ["Administrator (PMC)", "Government officials with elevated access", "Medium — data interpretation skills", "All citizen functions + user management, analytics, prediction accuracy review"],
    ["Unauthenticated Visitor", "Users not logged in", "Low", "View homepage, register, login"],
])
doc.add_paragraph(
    "The most important user class is the Citizen user, as the system is primarily designed to empower "
    "commuters with traffic insights and infrastructure reporting capabilities. The Administrator class "
    "is equally critical for municipal decision-making based on analytics."
)

doc.add_heading("2.4 Operating Environment", level=2)
add_tbl(doc, ["Component", "Specification"], [
    ["Client Platform", "Any modern web browser (Chrome, Firefox, Edge, Safari) on desktop or mobile"],
    ["Server Runtime", "Node.js v14+ (backend), Python 3.8+ (ML service)"],
    ["Database", "MongoDB 6.0+ (MongoDB Atlas cloud or local instance)"],
    ["Production Hosting", "Render.com (scalable cloud platform)"],
    ["Development Ports", "Frontend: 5173, Backend: 3001, ML Service: 3003"],
    ["Network", "Stable internet required for external API integrations"],
    ["Coexisting Software", "OpenWeatherMap, Google Maps, Cloudinary, OpenStreetMap APIs"],
])

doc.add_heading("2.5 Design and Implementation Constraints", level=2)
bp(doc, [
    "Backend must use Node.js with Express.js framework (MERN stack decision).",
    "Frontend must use React 18 with Vite as the build tool.",
    "All data models must use MongoDB with Mongoose ODM, supporting geospatial queries.",
    "Pothole detection model is pre-trained (MobileNet base) and served via Flask; no retraining in production.",
    "External APIs (OpenWeatherMap, Google Maps) have rate limits affecting data freshness.",
    "All user-uploaded images are stored via Cloudinary cloud service.",
    "JWT tokens stored in HTTP-only cookies with 7-day expiration.",
    "System designed specifically for Pune city routes and infrastructure.",
    "Traffic data collected every 2 hours in predefined time slots.",
    "Agile methodology with weekly sprints used for development.",
    "GitHub used for version control (separate backend and frontend repositories).",
])

doc.add_heading("2.6 User Documentation", level=2)
doc.add_paragraph("The following user documentation will be delivered with the software:")
bp(doc, [
    "This SRS document (Software Requirements Specification)",
    "B.Tech-14 Major Project Blackbook (comprehensive project documentation)",
    "README files in both frontend and backend repositories with setup instructions",
    "API documentation within the codebase and SRS Appendix D",
    "Inline code comments and modular architecture for developer reference",
])

doc.add_heading("2.7 Assumptions and Dependencies", level=2)
doc.add_paragraph("Assumptions:")
bp(doc, [
    "Citizens and authorities have access to an internet-enabled device.",
    "Real-time location and traffic data can be reliably accessed and updated.",
    "Users have modern web browsers with JavaScript enabled.",
    "The 5 monitored Pune routes remain the primary focus for automated data collection.",
])
doc.add_paragraph("Dependencies:")
bp(doc, [
    "OpenWeatherMap API — Weather data for traffic prediction scoring",
    "Google Maps Directions API — Real-time traffic-aware route durations",
    "OpenStreetMap — Base map tiles for Leaflet map rendering",
    "Cloudinary — Cloud-based image storage for complaints and pothole reports",
    "MongoDB Atlas — Primary data persistence with geospatial indexing",
    "TensorFlow/Keras — ML model inference for pothole detection (Flask service)",
])
doc.add_page_break()

# ============================================================
# 3. EXTERNAL INTERFACE REQUIREMENTS
# ============================================================
doc.add_heading("3. External Interface Requirements", level=1)

doc.add_heading("3.1 User Interfaces", level=2)
doc.add_paragraph("The system provides the following user interfaces:")
doc.add_paragraph("Citizen-facing interfaces:")
bp(doc, [
    "Home Page — Project overview with navigation to all features",
    "Login and Registration forms with input validation",
    "Traffic Prediction Page — Map-based route input with date/time slot selection",
    "Route Analysis Page — Dual-map (live traffic + reported issues) with traffic score display",
    "Pothole Registration — Image upload with ML verification and geolocation tagging",
    "Complaint Registration — Description, location, and image upload form",
    "Analytics Dashboard — Charts, heatmaps, and trend visualizations",
    "Calendar View — Date-based traffic data visualization",
    "Festival Analysis — Festival-impact traffic analysis",
    "Traffic Status — Real-time status updates",
    "Rewards/Redeem Page — Citizen contribution rewards",
])
doc.add_paragraph("Administrator interfaces:")
bp(doc, [
    "Admin Dashboard — Overview metrics and management controls",
    "Analytics Dashboard — Traffic trends and heatmap analysis",
    "Prediction Accuracy Dashboard — Predicted vs. actual score comparisons",
    "Pothole/Complaint Management — Review all reports with images and geolocation",
    "User Management — View users, promote to admin role",
])
doc.add_paragraph("All interfaces are responsive (desktop and mobile), use consistent navigation via a navbar component, "
    "use Chart.js for data visualization, SASS for styling, React-Leaflet for maps, and DOMPurify for content sanitization.")

doc.add_heading("3.2 Hardware Interfaces", level=2)
doc.add_paragraph(
    "No dedicated hardware interfaces are required. The system operates on standard cloud server "
    "infrastructure (Render.com) and MongoDB Atlas, eliminating dedicated server hardware needs. "
    "Client access requires any device with a modern web browser and internet connection."
)

doc.add_heading("3.3 Software Interfaces", level=2)
add_tbl(doc, ["Software", "Version/Type", "Purpose", "Data Exchange"], [
    ["OpenWeatherMap", "REST API (HTTPS)", "Weather data for Pune (18.5204, 73.8567)", "JSON — temperature, weather conditions, 5-day forecasts"],
    ["Google Maps Directions", "REST API (HTTPS)", "Real-time traffic durations", "JSON — route durations with/without traffic, delay ratios"],
    ["OpenStreetMap", "Tile Server (HTTPS)", "Map tiles for Leaflet rendering", "PNG tiles via HTTP"],
    ["Cloudinary", "SDK + REST API", "Image upload and storage", "multipart/form-data upload, URL retrieval"],
    ["MongoDB Atlas", "Mongoose ODM (TCP)", "Data persistence — 19 collections", "BSON documents via MongoDB protocol"],
    ["Flask ML Service", "REST API (HTTP:3003)", "Pothole image classification", "JSON — image classification result and confidence"],
])

doc.add_heading("3.4 Communications Interfaces", level=2)
doc.add_paragraph("The system uses the following communication protocols and standards:")
bp(doc, [
    "HTTP/HTTPS — Primary protocol for all web and API communication",
    "JSON — Standard data exchange format for all API requests and responses",
    "multipart/form-data — File upload protocol via Multer middleware for image submissions",
    "JWT (JSON Web Tokens) — Stateless authentication stored in HTTP-only cookies",
    "CORS — Cross-Origin Resource Sharing restricted to whitelisted origins (production URL + localhost:5173)",
    "bcrypt hashing — Password security with 10 salt rounds",
])
doc.add_page_break()

# ============================================================
# 4. SYSTEM FEATURES
# ============================================================
doc.add_heading("4. System Features", level=1)
doc.add_paragraph(
    "This section organizes the functional requirements by system features. Each feature includes "
    "a description and priority, stimulus/response sequences, and detailed functional requirements."
)

# 4.1 User Authentication
add_feature(doc, 1, "User Authentication and Management", "High",
    "The system provides user registration, login/logout, and role-based access control. "
    "Citizens self-register and administrators are promoted by existing admins. Authentication "
    "uses JWT tokens stored in HTTP-only cookies.",
    [
        "User navigates to Register page → fills form (username, email, password) → system creates account → returns profile data",
        "User navigates to Login page → enters email and password → system validates → issues JWT cookie → redirects to dashboard",
        "User clicks Logout → system clears JWT cookie → redirects to homepage",
        "Admin navigates to User Management → selects user → promotes to admin role",
    ],
    [
        ("REQ-AUTH-1", "The system shall provide a self-registration process requiring username, email, and password."),
        ("REQ-AUTH-2", "The system shall hash all passwords using bcrypt with 10 salt rounds before storage."),
        ("REQ-AUTH-3", "The system shall enforce unique constraints on both username and email fields."),
        ("REQ-AUTH-4", "The system shall return user profile data (excluding password) upon successful registration."),
        ("REQ-AUTH-5", "The system shall authenticate users with email and password, issuing a JWT token in an HTTP-only cookie with 7-day expiration."),
        ("REQ-AUTH-6", "The system shall implement role-based access control (RBAC) with User and Admin roles."),
        ("REQ-AUTH-7", "The system shall restrict admin routes (/admin/*) to users with isAdmin=true."),
        ("REQ-AUTH-8", "The system shall provide a logout endpoint that clears the JWT cookie."),
        ("REQ-AUTH-9", "The system shall allow admin users to promote regular users to admin role."),
    ]
)

# 4.2 Traffic Prediction Engine
add_feature(doc, 2, "Traffic Prediction Engine", "High",
    "The core analytical feature. Provides multi-factor traffic congestion prediction using a "
    "weighted scoring algorithm (0-100). Incorporates historical data, real-time weather, events, "
    "construction, hotspots, potholes, complaints, Google Traffic delay ratios, and time-of-day multipliers. "
    "Also provides smart time suggestions by comparing adjacent time slots.",
    [
        "User selects route, date, and time slot on prediction page → clicks 'Predict' → system fetches weather/traffic data from APIs → calculates weighted score → returns prediction with breakdown",
        "User requests time suggestions → system evaluates +/-2 adjacent time slots → returns ranked list with best travel time highlighted",
        "System caches weather data for 15 minutes to avoid redundant API calls",
    ],
    [
        ("REQ-PRED-1", "The system shall predict traffic congestion for a given route, date, and time slot using a weighted multi-factor scoring algorithm producing a score from 0 to 100."),
        ("REQ-PRED-2", "The system shall use the following weighted factors: Historical Average (W=25), Weather Impact (W=15), Construction Zones (W=15), Events (W=20), Hotspots (W=12), Potholes with decay (W=8), Complaints (W=5)."),
        ("REQ-PRED-3", "The system shall apply time-of-day multipliers: Morning Rush 08-10 (1.8x), Evening Rush 18-20 (1.8x), Afternoon Peak 12-14 (1.5x), Pre-Rush 16-18 (1.6x), Off-Peak 00-06 (0.5x)."),
        ("REQ-PRED-4", "The system shall integrate real-time weather data from OpenWeatherMap API with impact scores: Thunderstorm (+40), Rain (+35), Snow (+30), Fog (+25), Drizzle (+20), Clear (0)."),
        ("REQ-PRED-5", "The system shall fetch real-time traffic delay ratios from Google Maps Directions API."),
        ("REQ-PRED-6", "The system shall use the Haversine formula to calculate proximity of obstacles within a 200m radius."),
        ("REQ-PRED-7", "The system shall classify traffic levels: Low (0-25), Moderate (26-50), High (51-75), Severe (76-100)."),
        ("REQ-PRED-8", "The system shall cache weather data for 15 minutes to reduce external API calls."),
        ("REQ-PRED-9", "The system shall return a detailed breakdown of contributing factors with each prediction."),
        ("REQ-PRED-10", "The system shall compute predictions for +/-2 adjacent time slots and recommend the least congested alternatives."),
    ]
)

# 4.3 Pothole Detection
add_feature(doc, 3, "Pothole Detection (Machine Learning)", "High",
    "AI-powered pothole detection using a MobileNet deep learning model served via Flask microservice. "
    "Citizens upload road images which are classified as pothole or no-pothole. MobileNet was selected "
    "over ResNet-50 based on comparative performance analysis for deployment efficiency.",
    [
        "User navigates to Pothole page → captures/selects road image → enters location → submits form",
        "System uploads image to Cloudinary → sends to Flask ML service on port 3003 → preprocesses to 224x224 → runs MobileNet inference → returns classification result",
        "Result with geolocation stored in database → visible on admin dashboard and map",
    ],
    [
        ("REQ-POT-1", "The system shall accept uploaded road images (JPEG/PNG) via multipart/form-data."),
        ("REQ-POT-2", "The system shall process images through a Flask ML microservice on Port 3003."),
        ("REQ-POT-3", "The system shall preprocess images by resizing to 224x224 pixels and normalizing pixel values."),
        ("REQ-POT-4", "The system shall classify images using a pre-trained MobileNet model (pothole_mobnet_base.h5) with binary output: pothole (>=0.5) or no_pothole (<0.5)."),
        ("REQ-POT-5", "The system shall store detection results with geolocation coordinates in the Images collection."),
        ("REQ-POT-6", "The system shall upload submitted images to Cloudinary for persistent cloud storage."),
    ]
)

# 4.4 Complaint Management
add_feature(doc, 4, "Citizen Complaint Management", "High",
    "Enables citizens to report road-related issues through a structured complaint form with "
    "optional image uploads. Addresses the project objective of providing a platform for issue reporting.",
    [
        "User navigates to Complaint page → fills description and location → optionally uploads image → submits",
        "System stores complaint with user ID, uploads image to Cloudinary → complaint visible to admins",
        "Admin views all complaints on dashboard with location, timestamp, description, and image",
    ],
    [
        ("REQ-CMP-1", "The system shall allow authenticated citizens to submit complaints with description, location, and optional image upload."),
        ("REQ-CMP-2", "The system shall store complaint images on Cloudinary with UUID-based naming."),
        ("REQ-CMP-3", "The system shall associate each complaint with the submitting user's ID."),
        ("REQ-CMP-4", "The system shall provide an endpoint to retrieve all submitted complaints."),
        ("REQ-CMP-5", "The system shall display complaints with location, timestamp, description, and image on the admin panel."),
    ]
)

# 4.5 Event and Infrastructure Management
add_feature(doc, 5, "Event and Infrastructure Management", "Medium",
    "Manages traffic-affecting events and urban infrastructure data (10 facility types). Events "
    "are factored into traffic prediction. Infrastructure data includes schools with opening/closing "
    "times that contribute to peak traffic levels near those areas.",
    [
        "User/Admin creates event with name, location, time, impact → event stored and factored into predictions",
        "Admin manages facility data (schools, hospitals, etc.) via CRUD endpoints",
        "Prediction engine queries nearby facilities within route proximity using Haversine formula",
    ],
    [
        ("REQ-INF-1", "The system shall allow creation of events with name, description, location, start/end time, and expected impact."),
        ("REQ-INF-2", "The system shall provide endpoints to retrieve all registered events."),
        ("REQ-INF-3", "The system shall factor active events into traffic prediction scoring (W=20)."),
        ("REQ-INF-4", "The system shall manage 10 types of urban infrastructure: schools, hospitals, hotels, malls, banquet halls, gardens, parking buildings, construction projects, diversions, and metro stations."),
        ("REQ-INF-5", "Each facility record shall include geolocation data (latitude, longitude) and name."),
        ("REQ-INF-6", "The system shall provide CRUD endpoints for each facility type."),
        ("REQ-INF-7", "The system shall provide time-based relevance queries for schools (based on opening/closing timings)."),
        ("REQ-INF-8", "Construction projects and diversions shall include start/end dates and impact information."),
    ]
)

# 4.6 Traffic Analytics
add_feature(doc, 6, "Traffic Analytics and Reporting", "High",
    "Comprehensive analytics engine providing traffic trends, heatmaps, root-cause analysis, "
    "route comparisons, and AI-generated recommendations. Supports festival-specific and "
    "calendar-based analysis views.",
    [
        "User navigates to Analytics page → selects route and date range → system queries historical data → renders charts",
        "Admin views heatmap → system generates time-slots-vs-days matrix → displays color-coded visualization",
        "System generates AI recommendations based on root-cause factor breakdown",
    ],
    [
        ("REQ-ANA-1", "The system shall provide traffic trend analysis over configurable N days for a given route."),
        ("REQ-ANA-2", "The system shall generate heatmap data (time slots vs. days matrix) for visualization."),
        ("REQ-ANA-3", "The system shall provide root-cause breakdown showing factor-by-factor congestion contribution."),
        ("REQ-ANA-4", "The system shall provide route comparison data for a given day."),
        ("REQ-ANA-5", "The system shall generate AI-powered improvement recommendations from root-cause analysis."),
        ("REQ-ANA-6", "The system shall provide a list of all available monitored routes."),
        ("REQ-ANA-7", "The system shall provide festival-specific traffic analysis."),
        ("REQ-ANA-8", "The system shall provide calendar-based traffic data views."),
        ("REQ-ANA-9", "The system shall provide last-four-week weekday trend analysis."),
    ]
)

# 4.7 Automated Data Collection
add_feature(doc, 7, "Automated Data Collection", "High",
    "Scheduled background service that automatically collects traffic data for 5 monitored "
    "Pune routes every 2 hours using node-cron scheduler. Generates realistic scores with "
    "detailed breakdowns.",
    [
        "System scheduler triggers every 2 hours → collects data for each monitored route → generates traffic scores and breakdowns → stores in PathInfo collection",
        "Data includes pathId, date, timeSlot, score, trafficLevel, and factor breakdown",
    ],
    [
        ("REQ-DAT-1", "The system shall automatically collect traffic data for 5 monitored routes every 2 hours via node-cron."),
        ("REQ-DAT-2", "Monitored routes shall include: Hinjewadi-Swargate, Hadapsar-Shivajinagar, Kothrud-Shivajinagar, Swargate-Katraj, and additional configured routes."),
        ("REQ-DAT-3", "The system shall generate realistic traffic scores based on time patterns and route characteristics."),
        ("REQ-DAT-4", "The system shall create detailed factor breakdowns for each collected data point."),
        ("REQ-DAT-5", "The system shall store data in PathInfo collection with pathId, date, timeSlot, score, trafficLevel, and breakdown."),
    ]
)

# 4.8 Map Visualization
add_feature(doc, 8, "Map Visualization and Location Services", "High",
    "Dual-map interface using Leaflet (React-Leaflet) and Google Maps. One map displays live "
    "traffic, another shows reported complaints, potholes, events, and nearby institutions. "
    "Supports route visualization, geosearch, and proximity detection.",
    [
        "User opens analysis page → map loads with markers for hotspots, potholes, complaints, events",
        "User searches for a location → Leaflet GeoSearch returns results → map centers on location",
        "User selects route → system calculates nearby facilities using Haversine within 200m radius",
    ],
    [
        ("REQ-MAP-1", "The system shall display interactive maps using Leaflet with React-Leaflet."),
        ("REQ-MAP-2", "The system shall support route visualization via Leaflet Routing Machine."),
        ("REQ-MAP-3", "The system shall provide location search via Leaflet GeoSearch."),
        ("REQ-MAP-4", "The system shall integrate Google Maps API for additional route capabilities."),
        ("REQ-MAP-5", "The system shall detect nearby facilities along a route using Haversine formula within 200m radius."),
        ("REQ-MAP-6", "The system shall display hotspots, potholes, complaints, events, and institutions as markers on the map."),
    ]
)

# 4.9 Admin Dashboard
add_feature(doc, 9, "Admin Dashboard", "High",
    "Dedicated administrative panel accessible only to users with admin role. Provides "
    "management views for all system data including pothole reports, complaints, users, "
    "analytics, and prediction accuracy tracking.",
    [
        "Admin logs in → redirected to admin dashboard with overview metrics",
        "Admin navigates to Pothole Reports → views all ML detection results with images and geolocation",
        "Admin navigates to Users → views user list → can promote selected user to admin",
        "Admin views Prediction Accuracy → sees predicted vs. actual scores over time",
    ],
    [
        ("REQ-ADM-1", "The system shall provide an admin dashboard accessible only to users with isAdmin=true."),
        ("REQ-ADM-2", "The dashboard shall display all pothole detection reports with images and geolocation."),
        ("REQ-ADM-3", "The dashboard shall display all citizen complaints with details and images."),
        ("REQ-ADM-4", "The dashboard shall provide user management with role promotion capability."),
        ("REQ-ADM-5", "The dashboard shall provide analytics views with traffic trends and heatmaps."),
        ("REQ-ADM-6", "The dashboard shall provide prediction accuracy tracking (predicted vs. actual scores)."),
    ]
)

# 4.10 Prediction Accuracy
add_feature(doc, 10, "Prediction Accuracy and Logging", "Medium",
    "Tracks prediction accuracy by logging each prediction and backfilling actual scores when "
    "real data arrives. Calculates accuracy percentage for system performance monitoring.",
    [
        "System generates prediction → logs predictedScore in PredictionLog",
        "Actual data arrives → system backfills actualScore → calculates accuracy percentage → marks as verified",
    ],
    [
        ("REQ-LOG-1", "The system shall log each prediction with pathId, date, timeSlot, and predictedScore."),
        ("REQ-LOG-2", "The system shall backfill actual scores into prediction logs when real data arrives."),
        ("REQ-LOG-3", "The system shall calculate accuracy percentage by comparing predicted vs. actual scores."),
        ("REQ-LOG-4", "The system shall mark prediction logs as verified after backfill."),
    ]
)

# 4.11 Citizen Rewards
add_feature(doc, 11, "Citizen Rewards", "Low",
    "Gamification system that rewards citizens for their contributions to the platform "
    "(reporting potholes, submitting complaints). Citizens accumulate scores viewable on a rewards page.",
    [
        "User reports pothole or complaint → system increments citizen score",
        "User navigates to Rewards page → views accumulated score",
    ],
    [
        ("REQ-RWD-1", "The system shall maintain a citizen score for each registered user based on contributions."),
        ("REQ-RWD-2", "The system shall provide a redemption/rewards page for viewing accumulated score."),
    ]
)
doc.add_page_break()

# ============================================================
# 5. OTHER NONFUNCTIONAL REQUIREMENTS
# ============================================================
doc.add_heading("5. Other Nonfunctional Requirements", level=1)

doc.add_heading("5.1 Performance Requirements", level=2)
bp(doc, [
    "API response times shall be less than 500ms for standard CRUD operations under normal load.",
    "Traffic prediction computation shall complete within 3 seconds, including external API calls.",
    "Pothole detection ML inference shall complete within 5 seconds per image.",
    "Weather data shall be cached for 15 minutes to minimize external API latency.",
    "The system shall support at least 100 concurrent users during normal operations.",
    "Automated data collection shall process all 5 monitored routes within each 2-hour cycle.",
    "The system shall use a stateless API design to enable horizontal scaling.",
    "node-cache shall be used for frequently accessed data to reduce database load.",
])

doc.add_heading("5.2 Safety Requirements", level=2)
doc.add_paragraph(
    "The system provides traffic predictions and analytics for informational purposes only. "
    "Traffic predictions are statistical estimates and should not be used as the sole basis for "
    "safety-critical decisions. The system includes appropriate disclaimers that actual traffic "
    "conditions may vary from predictions."
)

doc.add_heading("5.3 Security Requirements", level=2)
bp(doc, [
    "All passwords shall be hashed using bcrypt with 10 salt rounds.",
    "Authentication shall use JWT tokens in HTTP-only cookies with 7-day expiration.",
    "Role-based access control (RBAC) shall be enforced with User and Admin roles.",
    "All API inputs shall be validated using express-validator middleware.",
    "User-generated content shall be sanitized on the frontend using DOMPurify.",
    "CORS shall be restricted to whitelisted origins only.",
    "All environment secrets shall be stored in .env files, never hardcoded in source code.",
    "Sensitive data shall be encrypted and securely stored on MongoDB Atlas.",
])

doc.add_heading("5.4 Software Quality Attributes", level=2)
doc.add_paragraph("Reliability:")
bp(doc, [
    "The scheduler shall run reliably every 2 hours with error handling.",
    "MongoDB connection retry logic shall be implemented on failure.",
    "The system shall degrade gracefully when external APIs are unavailable.",
])
doc.add_paragraph("Usability:")
bp(doc, [
    "The application shall be responsive on desktop and mobile browsers.",
    "Intuitive map-based route selection and visualization shall be provided.",
    "Chart.js charts and heatmaps shall enable quick traffic pattern comprehension.",
])
doc.add_paragraph("Maintainability:")
bp(doc, [
    "The codebase shall follow a modular MVC pattern (Models, Controllers, Routes).",
    "ES Modules shall be used throughout (type: 'module' in package.json).",
    "All configuration shall be externalized via environment variables.",
    "GitHub shall be used for version control with separate backend/frontend repositories.",
])

doc.add_heading("5.5 Business Rules", level=2)
bp(doc, [
    "Only authenticated users can submit complaints, pothole reports, or access predictions.",
    "Only admin users (isAdmin=true) can access the admin dashboard and management functions.",
    "Only admin users can promote other users to admin role.",
    "Traffic data collection runs automatically every 2 hours regardless of user activity.",
    "Prediction accuracy is calculated only after actual data is backfilled for a time period.",
])
doc.add_page_break()

# ============================================================
# 6. OTHER REQUIREMENTS
# ============================================================
doc.add_heading("6. Other Requirements", level=1)
doc.add_paragraph(
    "The system shall be developed using Agile methodology with weekly sprints. "
    "GitHub shall be used for source control. The system is designed specifically for "
    "Pune city routes and infrastructure, with data models supporting geospatial queries "
    "for location-based filtering. The database must store 19 distinct collections for "
    "traffic data, user reports, infrastructure records, and analytics."
)
doc.add_page_break()

# ============================================================
# APPENDIX A: GLOSSARY
# ============================================================
doc.add_heading("Appendix A: Glossary", level=1)
add_tbl(doc, ["Term", "Definition"], [
    ["PMC", "Pune Municipal Corporation — the sponsoring government body"],
    ["JWT", "JSON Web Token — used for stateless authentication"],
    ["RBAC", "Role-Based Access Control"],
    ["API", "Application Programming Interface"],
    ["REST", "Representational State Transfer"],
    ["ML", "Machine Learning"],
    ["MobileNet", "Lightweight deep learning architecture for image classification"],
    ["CRUD", "Create, Read, Update, Delete"],
    ["PathInfo", "Data record storing traffic score, time slot, date, and route info"],
    ["Traffic Score", "Numerical value (0-100) representing congestion severity"],
    ["Time Slot", "2-hour window (e.g., '08-10') for data collection and prediction"],
    ["Hotspot", "Known traffic congestion point with geographic coordinates"],
    ["Haversine", "Formula to calculate distance between two GPS coordinates"],
    ["CORS", "Cross-Origin Resource Sharing"],
    ["Agile", "Iterative development methodology with weekly sprints"],
    ["SASS", "CSS preprocessor for frontend styling"],
])
doc.add_paragraph()

# ============================================================
# APPENDIX B: ANALYSIS MODELS
# ============================================================
doc.add_heading("Appendix B: Analysis Models", level=1)
doc.add_paragraph("System Architecture (Three-Tier):")
doc.add_paragraph(
    "The system follows a three-tier architecture pattern:\n"
    "Tier 1 (Presentation): React + Vite SPA with React-Leaflet maps and Chart.js visualizations\n"
    "Tier 2 (Application): Node.js Express REST API (Port 3001) + Flask ML Service (Port 3003)\n"
    "Tier 3 (Data): MongoDB Atlas with 19 collections and geospatial indexing"
)
doc.add_paragraph()
doc.add_paragraph("Data Flow — Traffic Prediction:")
doc.add_paragraph(
    "1. User selects route/date/time → Frontend sends GET /api/path-info/predictTraffic\n"
    "2. Backend queries: Historical PathInfo, Events, Constructions, Hotspots, Potholes, Complaints\n"
    "3. Backend calls: OpenWeatherMap API (weather), Google Maps API (traffic delay)\n"
    "4. Weighted scoring: W_EVENT(20) + W_HISTORICAL(25) + W_WEATHER(15) + W_CONSTRUCTION(15) + W_HOTSPOT(12) + W_POTHOLE(8) + W_COMPLAINT(5)\n"
    "5. Time-of-day multiplier applied → Final score (0-100) → Traffic level classification\n"
    "6. Response with score, breakdown, recommendations returned to frontend"
)
doc.add_paragraph()
doc.add_paragraph("Data Flow — Pothole Detection:")
doc.add_paragraph(
    "1. User uploads image → Backend receives via Multer middleware\n"
    "2. Image uploaded to Cloudinary → URL stored\n"
    "3. Image forwarded to Flask service (Port 3003)\n"
    "4. Flask preprocesses (resize 224x224, normalize) → MobileNet inference\n"
    "5. Binary classification: pothole (>=0.5) / no_pothole (<0.5)\n"
    "6. Result stored in Images collection with geolocation"
)
doc.add_page_break()

# ============================================================
# APPENDIX C: TECHNOLOGY STACK
# ============================================================
doc.add_heading("Appendix C: Technology Stack", level=1)
doc.add_paragraph("Backend (Node.js Express Server):")
add_tbl(doc, ["Category", "Technology", "Version"], [
    ["Runtime", "Node.js", "v14+"],
    ["Framework", "Express.js", "4.21.0"],
    ["Database", "MongoDB (Mongoose ODM)", "8.20.1"],
    ["Authentication", "jsonwebtoken", "9.0.2"],
    ["Password Hashing", "bcryptjs", "3.0.3"],
    ["File Upload", "Multer + Cloudinary", "1.4.5 / 1.41.3"],
    ["Validation", "express-validator", "7.3.1"],
    ["HTTP Client", "Axios", "1.13.2"],
    ["Scheduling", "node-cron", "3.0.3"],
    ["Caching", "node-cache", "5.1.2"],
])
doc.add_paragraph()
doc.add_paragraph("ML Microservice (Flask):")
add_tbl(doc, ["Category", "Technology"], [
    ["Framework", "Flask"],
    ["ML Framework", "TensorFlow / Keras"],
    ["Image Processing", "Pillow, NumPy"],
    ["Model", "MobileNet Base (pothole_mobnet_base.h5, 224x224)"],
])
doc.add_paragraph()
doc.add_paragraph("Frontend (React + Vite):")
add_tbl(doc, ["Category", "Technology", "Version"], [
    ["Framework", "React", "18.3.1"],
    ["Build Tool", "Vite", "5.4.1"],
    ["Routing", "React Router DOM", "6.26.2"],
    ["Maps", "Leaflet + React-Leaflet", "1.9.4 / 4.2.1"],
    ["Maps (Google)", "@react-google-maps/api", "2.20.3"],
    ["Charts", "Chart.js + react-chartjs-2", "4.4.7 / 5.2.0"],
    ["Styling", "SASS", "1.79.4"],
])
doc.add_page_break()

# ============================================================
# APPENDIX D: API SPECIFICATION
# ============================================================
doc.add_heading("Appendix D: API Specification", level=1)

api_groups = [
    ("Authentication", [
        ("POST", "/api/auth/register", "Register new citizen user", "No"),
        ("POST", "/api/auth/login", "Login and receive JWT cookie", "No"),
        ("POST", "/api/auth/logout", "Clear JWT cookie", "Yes"),
    ]),
    ("Traffic Prediction", [
        ("GET", "/api/path-info/predictTraffic", "Multi-factor traffic prediction", "Yes"),
        ("GET", "/api/path-info/getTimeSuggestions", "Smart time suggestions", "Yes"),
        ("POST", "/api/path-info", "Add traffic path data", "Yes"),
        ("GET", "/api/path-info/getCalendarData", "Calendar-based traffic data", "Yes"),
        ("GET", "/api/path-info/getFestivalData", "Festival traffic analysis", "Yes"),
        ("GET", "/api/path-info/getLastFourWeek", "Last 4 weeks trends", "Yes"),
    ]),
    ("Analytics", [
        ("GET", "/api/analytics/trends", "Traffic trends over N days", "Yes"),
        ("GET", "/api/analytics/heatmap", "Heatmap matrix data", "Yes"),
        ("GET", "/api/analytics/breakdown", "Root-cause breakdown", "Yes"),
        ("GET", "/api/analytics/recommendations", "AI recommendations", "Yes"),
        ("GET", "/api/analytics/comparison", "Route comparison", "Yes"),
        ("GET", "/api/analytics/routes", "Available routes list", "Yes"),
    ]),
    ("Complaints & Potholes", [
        ("POST", "/api/complaint", "Submit complaint with image", "Yes"),
        ("GET", "/api/complaint/getComplaintData", "Get all complaints", "Yes"),
        ("POST", "/api/model", "Upload for pothole detection", "Yes"),
        ("GET", "/api/model/getPotholeData", "Get pothole results", "Yes"),
    ]),
]
for grp_name, endpoints in api_groups:
    doc.add_heading(grp_name, level=2)
    add_tbl(doc, ["Method", "Endpoint", "Description", "Auth"], [list(e) for e in endpoints])
    doc.add_paragraph()
doc.add_page_break()

# ============================================================
# APPENDIX E: DATA DICTIONARY
# ============================================================
doc.add_heading("Appendix E: Data Dictionary", level=1)
doc.add_paragraph("Key MongoDB Collections (19 total):")
add_tbl(doc, ["Collection", "Key Fields", "Purpose"], [
    ["users", "username, email, password (hashed), isAdmin", "User accounts with RBAC"],
    ["pathinfos", "pathId, date, timeSlot, score, breakdown, trafficLevel", "Historical traffic records"],
    ["predictionlogs", "pathId, date, timeSlot, predictedScore, actualScore, accuracy, verified", "Prediction tracking"],
    ["complaints", "description, location, imageUrl, userId", "Citizen complaints"],
    ["images", "imageUrl, location, result, confidence", "Pothole detection results"],
    ["events", "name, description, location, startTime, endTime, expectedImpact", "Traffic events"],
    ["constructions", "name, location, startDate, endDate, impact", "Construction projects"],
    ["diversions", "name, location, reason", "Traffic diversions"],
    ["schools", "name, latitude, longitude, timing", "Schools with timings"],
    ["hospitals", "name, latitude, longitude", "Healthcare facilities"],
    ["hotels", "name, latitude, longitude", "Hospitality venues"],
    ["malls", "name, latitude, longitude", "Shopping centers"],
    ["banquethalls", "name, latitude, longitude, capacity", "Event venues"],
    ["gardens", "name, latitude, longitude", "Public parks"],
    ["parkingbuildings", "name, latitude, longitude", "Parking facilities"],
    ["hotspotlocations", "name, latitude, longitude", "Congestion hotspots"],
    ["trafficstatuses", "location, status, timestamp", "Real-time updates"],
    ["metrostations", "name, latitude, longitude, line", "Metro stations"],
])

# ============ END ============
doc.add_paragraph(); doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— End of SRS Document —"); r.bold = True; r.font.size = Pt(14)

# Save
output = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SRS_Traffic_Prediction_IEEE.docx")
doc.save(output)
print(f"Done: {output}")
