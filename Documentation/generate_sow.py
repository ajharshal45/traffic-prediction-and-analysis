from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Title
title = doc.add_heading('Statement of Work (SOW)', level=1)
for run in title.runs:
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 0, 0)

# Section 1: Project Title
doc.add_heading('1. Project Title', level=2)
doc.add_paragraph('Data Driven Pune Municipal Corporation (PMC) Traffic Analysis System')

# Section 2: Vision
doc.add_heading('2. Vision', level=2)
doc.add_paragraph(
    'The vision of this project is to develop a comprehensive, data-driven, and AI-powered traffic analysis '
    'and prediction system for Pune Municipal Corporation (PMC) that integrates real-time traffic monitoring, '
    'predictive analytics, pothole detection, citizen complaint management, and urban infrastructure tracking '
    'on a single unified web platform. The system aims to empower both citizens and municipal authorities with '
    'actionable insights for smarter commuting decisions and proactive urban traffic management, ultimately '
    'contributing to safer roads, reduced congestion, and improved quality of urban mobility in Pune city.'
)

# Section 3: Problem Statement
doc.add_heading('3. Problem Statement', level=2)
doc.add_paragraph(
    'In rapidly urbanizing cities like Pune, road congestion, poorly maintained infrastructure, and inefficient '
    'management of public spaces are recurring challenges that significantly affect daily commuting. The lack of '
    'accessible systems for reporting infrastructure issues, such as potholes or traffic hotspots, leads to delays '
    'in repairs and contributes to unsafe travel conditions. Existing road issue reporting and traffic information '
    'systems are fragmented and lack centralized access, making them inconvenient for citizens. Route analysis is '
    'limited to basic congestion data, without consideration of local factors like school traffic, potholes, '
    'construction zones, events, and hotspots, reducing the accuracy and usefulness of travel insights. Without '
    'a centralized system to report, analyze, and address these factors, both commuters and city authorities '
    'struggle with unpredictable traffic patterns and infrastructure strain, impacting the quality of life and '
    'urban mobility.'
)

# Section 4: Goals
doc.add_heading('4. Goals', level=2)
doc.add_paragraph('The primary goals of this project are:')
goals = [
    'To provide a centralized digital platform for real-time traffic analysis, prediction, and route planning for Pune city roads',
    'To enable AI-powered pothole detection using deep learning (MobileNet) to expedite road maintenance',
    'To facilitate citizen engagement through a complaint management and reporting system with geolocation and image support',
    'To deliver multi-factor traffic congestion prediction incorporating historical data, weather, events, construction, hotspots, and nearby facilities',
    'To support Pune Municipal Corporation in data-driven decision-making for urban traffic management and infrastructure maintenance',
    'To improve road safety and commuter experience through actionable analytics, smart time suggestions, and AI-generated recommendations',
]
for g in goals:
    doc.add_paragraph(f'- {g}', style='List Bullet')

# Section 5: Objectives
doc.add_heading('5. Objectives', level=2)
objectives = [
    'Design and develop a full-stack web application using React (Vite) frontend, Node.js (Express) backend, and Flask ML microservice',
    'Implement a weighted multi-factor traffic prediction engine using historical data, weather (OpenWeatherMap API), Google Traffic data, events, construction zones, hotspots, potholes, complaints, and nearby facilities',
    'Deploy a MobileNet-based deep learning model for automatic pothole detection from citizen-uploaded road images',
    'Build an interactive map interface using Leaflet and Google Maps for route visualization, geosearch, and infrastructure mapping',
    'Develop analytics dashboards with traffic trends, heatmaps, root-cause breakdowns, route comparisons, and AI recommendations',
    'Implement automated traffic data collection via scheduled services (node-cron) every 2 hours for 5 monitored Pune routes',
    'Enable secure user authentication with JWT tokens, role-based access control (RBAC), and an admin management panel',
    'Provide smart time suggestions by comparing adjacent time slots to recommend optimal travel times',
]
for obj in objectives:
    doc.add_paragraph(f'- {obj}', style='List Bullet')

# Section 6: Scope of Work
doc.add_heading('6. Scope of Work', level=2)
doc.add_paragraph(
    'The scope of this project includes requirement analysis, system design, development, testing, deployment, '
    'and documentation of the Data Driven PMC Traffic Analysis System. The project encompasses the following areas:'
)
scope_items = [
    'Backend Development: RESTful API server using Node.js Express with 20 route files, 22 controllers, and 19 Mongoose database models',
    'Frontend Development: Single Page Application (SPA) using React 18 with Vite, featuring interactive maps (Leaflet, Google Maps), data visualization (Chart.js), and responsive design',
    'Machine Learning Service: Flask-based microservice for pothole detection using a pre-trained MobileNet model',
    'Database Design: MongoDB database (trafficanalysis) with 19 collections covering users, traffic data, complaints, events, facilities, and urban infrastructure',
    'External API Integration: OpenWeatherMap API for weather data, Google Maps Directions API for real-time traffic, and Cloudinary for image storage',
    'Automated Data Pipeline: Scheduled data collection service for 5 monitored Pune routes every 2 hours with prediction accuracy tracking',
    'Urban Infrastructure Tracking: Management of schools, hospitals, hotels, malls, banquet halls, gardens, parking buildings, construction projects, diversions, and metro stations',
    'Testing and Deployment: Application testing, deployment on Render.com, and comprehensive project documentation',
]
for s in scope_items:
    doc.add_paragraph(f'- {s}', style='List Bullet')

# Section 7: Deliverables
doc.add_heading('7. Deliverables', level=2)
deliverables = [
    'Software Requirement Specification (SRS) document',
    'Statement of Work (SOW) document',
    'System Design Documents (architecture diagrams, database schema, API specification)',
    'Fully functional web application (React frontend + Express backend + Flask ML service)',
    'Trained MobileNet model for pothole detection (pothole_mobnet_base.h5)',
    'Automated data collection and scheduling service',
    'Admin dashboard for system management',
    'Test reports, user manual, and project blackbook',
]
for d in deliverables:
    doc.add_paragraph(f'- {d}', style='List Bullet')

# Section 8: Stakeholders
doc.add_heading('8. Stakeholders', level=2)
stakeholders = [
    'Citizens (Regular Users) \u2014 General public who view traffic predictions, submit complaints, report potholes, and access analytics',
    'Pune Municipal Corporation (PMC) \u2014 Sponsor organization; uses the system for traffic management and infrastructure oversight',
    'System Administrator \u2014 Manages users, reviews complaints and pothole reports, oversees analytics and prediction accuracy',
    'Project Guide (Dr. Radhika Kulkarni) \u2014 Academic guide providing mentorship and project evaluation',
    'Development Team (Group A-14) \u2014 Ajinkya Walunj, Akash Bhandari, Anish Sagri, Annany Dev Singh',
    'IEEE Pune Chapter Smart City Section \u2014 Dr. Amar Buchade and Dr. Tushar Mote providing support and resources',
]
for s in stakeholders:
    doc.add_paragraph(f'- {s}', style='List Bullet')

# Section 9: Assumptions and Constraints
doc.add_heading('9. Assumptions and Constraints', level=2)

p_assumptions = doc.add_paragraph()
run = p_assumptions.add_run('Assumptions:')
run.bold = True

assumptions = [
    'Users have access to modern web browsers with JavaScript enabled and stable internet connectivity.',
    'The MongoDB database server (local or MongoDB Atlas cloud) is available and accessible.',
    'External APIs (OpenWeatherMap, Google Maps Directions, Cloudinary) remain available and functional.',
    'The 5 monitored Pune routes remain the primary focus for automated data collection.',
    'Real-time location and traffic data can be reliably accessed and updated on the portal.',
    'Citizens and municipal authorities will have access to an internet-enabled device for using the platform.',
]
for a in assumptions:
    doc.add_paragraph(f'- {a}', style='List Bullet')

p_constraints = doc.add_paragraph()
run = p_constraints.add_run('Constraints:')
run.bold = True

constraints = [
    'Project timeline is limited to the academic year 2024\u20132025.',
    'API rate limits on OpenWeatherMap and Google Maps may affect data freshness and collection frequency.',
    'The ML model (MobileNet) is pre-trained and not retrained in production; accuracy depends on training data quality.',
    'Geographic scope is limited to Pune city routes.',
    'High traffic volumes and concurrent users could impact system performance during peak hours.',
    'All user-uploaded images must be stored via Cloudinary (no local persistent storage).',
    'Integration of ML-based image verification for pothole reporting adds system complexity.',
]
for c in constraints:
    doc.add_paragraph(f'- {c}', style='List Bullet')

# Section 10: Conclusion
doc.add_heading('10. Conclusion', level=2)
doc.add_paragraph(
    'This Statement of Work outlines the vision, goals, and objectives of developing the Data Driven Pune '
    'Municipal Corporation (PMC) Traffic Analysis System. The project addresses real-world challenges in urban '
    'traffic management, road safety, and infrastructure reporting by providing a comprehensive, AI-powered web '
    'platform. Through multi-factor traffic prediction, automated data collection, pothole detection using deep '
    'learning, citizen complaint management, and rich analytics dashboards, the system aims to transform how '
    "Pune's traffic data is collected, analyzed, and acted upon \u2014 ultimately contributing to smarter, safer, "
    'and more efficient urban mobility for all stakeholders.'
)

# Save
doc.save('SOW_PMC_Traffic_Analysis.docx')
print('SOW document created successfully!')
